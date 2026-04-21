"""
Amazon転売リサーチツール メインスクリプト

使い方:
    python main.py 0421        # 月日4桁 → ~/Downloads/KeepaExport-2026-04-21.csv を使用
    python main.py path/to/file.csv  # フルパスも可

出力:
    ~/Desktop/せどりリサーチ結果_YYYYMMDD.xlsx
"""

import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from keepa_parser import parse_keepa_csv, KeepaProduct
from price_fetcher import fetch_purchase_options_split, PurchaseOption
from profit_calculator import (
    calculate_profit, ProfitResult,
    JUDGMENT_PROFITABLE, JUDGMENT_MARGINAL, JUDGMENT_LOSS,
)
import config

JUDGMENT_NO_DATA = "データなし"

JUDGMENT_COLORS = {
    JUDGMENT_PROFITABLE: "E2EFDA",  # 緑
    JUDGMENT_MARGINAL:   "FFF2CC",  # 黄
    JUDGMENT_LOSS:       "FCE4D6",  # 赤
    JUDGMENT_NO_DATA:    "F2F2F2",  # グレー
}


def _resolve_csv_path(arg: str) -> Path:
    """4桁の月日（例: 0421）またはファイルパスを受け取りCSVのPathを返す"""
    if arg.isdigit() and len(arg) == 4:
        year = datetime.now().year
        month, day = arg[:2], arg[2:]
        return Path.home() / "Downloads" / f"KeepaExport-{year}-{month}-{day}.csv"
    return Path(arg)


def main():
    if len(sys.argv) < 2:
        print("使い方: python3 main.py 0421  （月日4桁）")
        print("        python3 main.py ~/Downloads/KeepaExport-2026-04-21.csv  （フルパスも可）")
        sys.exit(1)

    csv_path = _resolve_csv_path(sys.argv[1])
    if not csv_path.exists():
        print(f"ファイルが見つかりません: {csv_path}")
        sys.exit(1)

    print(f"Keepa CSVを読み込んでいます: {csv_path}")
    products = parse_keepa_csv(csv_path)
    print(f"  {len(products)} 商品を読み込みました")

    jan_products = [p for p in products if p.jan]
    print(f"  うちJANコードあり: {len(jan_products)} 商品\n")

    rows = []
    for i, product in enumerate(jan_products, 1):
        print(f"[{i}/{len(jan_products)}] {product.asin} / JAN:{product.jan}")
        print(f"  Amazon価格: ¥{product.amazon_price:,} | {product.title[:40]}")

        rakuten_list, yahoo_list = fetch_purchase_options_split(product.jan)

        rakuten_best = rakuten_list[0] if rakuten_list else None
        yahoo_best   = yahoo_list[0]   if yahoo_list   else None

        if not rakuten_best and not yahoo_best:
            print("  → 仕入れ候補なし（データなし）")
            rows.append(_make_row(i, product, None, None, JUDGMENT_NO_DATA, None, None))
            continue

        pr_rakuten = calculate_profit(
            amazon_price=product.amazon_price,
            purchase_total=rakuten_best.total,
            category=product.category or "",
        ) if rakuten_best else None

        pr_yahoo = calculate_profit(
            amazon_price=product.amazon_price,
            purchase_total=yahoo_best.total,
            category=product.category or "",
        ) if yahoo_best else None

        # 最安仕入先の判定（同額の場合は楽天を優先）
        candidates = [(opt, pr) for opt, pr in [(rakuten_best, pr_rakuten), (yahoo_best, pr_yahoo)] if opt]
        best_opt, best_pr = min(candidates, key=lambda x: (x[0].total, 0 if x[0].source == "rakuten" else 1))

        print(f"  楽天最安: {'¥' + f'{rakuten_best.total:,}' if rakuten_best else 'なし'}"
              f"  ヤフー最安: {'¥' + f'{yahoo_best.total:,}' if yahoo_best else 'なし'}")
        print(f"  粗利(楽天): {'¥' + f'{pr_rakuten.profit:,}' if pr_rakuten else '-'}"
              f"  粗利(ヤフー): {'¥' + f'{pr_yahoo.profit:,}' if pr_yahoo else '-'}"
              f"  → {best_pr.judgment}")

        rows.append(_make_row(i, product, rakuten_best, yahoo_best, best_pr.judgment, pr_rakuten, pr_yahoo))

    profitable = [r for r in rows if r["judgment"] == JUDGMENT_PROFITABLE]
    print(f"\n=== 完了: {len(profitable)} / {len(rows)} 商品が利益あり（粗利1000円以上） ===\n")

    output_path = save_results(rows)
    print(f"結果を保存しました: {output_path}")


def _make_row(no, product, rakuten_best, yahoo_best, judgment, pr_rakuten, pr_yahoo):
    return {
        "no": no,
        "product": product,
        "rakuten_best": rakuten_best,
        "yahoo_best": yahoo_best,
        "judgment": judgment,
        "pr_rakuten": pr_rakuten,
        "pr_yahoo": pr_yahoo,
    }


def save_results(all_rows: list) -> str:
    wb = openpyxl.Workbook()

    # 赤字・データなしを除外
    filtered = [r for r in all_rows if r["judgment"] not in (JUDGMENT_LOSS, JUDGMENT_NO_DATA)]

    ws1 = wb.active
    ws1.title = "全商品"
    _write_sheet(ws1, filtered)

    profitable = [r for r in filtered if r["judgment"] == JUDGMENT_PROFITABLE]
    ws2 = wb.create_sheet("利益あり商品")
    _write_sheet(ws2, profitable)

    date_str = datetime.now().strftime("%Y%m%d")
    desktop = Path.home() / "Desktop" / f"せどりリサーチ結果_{date_str}.xlsx"
    wb.save(desktop)
    return str(desktop)


HEADERS = [
    "判定", "No.", "商品名", "Amazon現在価格", "楽天最安値", "ヤフー最安値", "ショップ名",
    "ASIN", "EAN", "最安仕入先", "仕入最安値",
    "Amazon手数料", "FBA配送料",
    "粗利(楽天仕入)", "粗利(ヤフー仕入)",
    "楽天最安値ページURL", "ヤフー最安値ページURL",
]

LINK_FONT   = Font(color="0563C1", underline="single")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2E75B6")


def _write_sheet(ws, rows: list):
    # ヘッダー行
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # オートフィルター
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    for r, row in enumerate(rows, 2):
        p: KeepaProduct       = row["product"]
        rb: PurchaseOption    = row.get("rakuten_best")
        yb: PurchaseOption    = row.get("yahoo_best")
        pr_r: ProfitResult    = row.get("pr_rakuten")
        pr_y: ProfitResult    = row.get("pr_yahoo")
        judgment: str         = row.get("judgment", JUDGMENT_NO_DATA)

        # 最安仕入先（同額の場合は楽天を優先）
        if rb and yb:
            best_source = "楽天" if rb.total <= yb.total else "ヤフー"
            best_price  = min(rb.total, yb.total)
        elif rb:
            best_source, best_price = "楽天", rb.total
        elif yb:
            best_source, best_price = "ヤフー", yb.total
        else:
            best_source, best_price = "", ""

        # Amazon手数料・FBA料金（楽天優先、なければヤフー）
        ref_fee = (pr_r or pr_y).referral_fee if (pr_r or pr_y) else ""
        fba_fee = (pr_r or pr_y).fba_fee      if (pr_r or pr_y) else ""

        best_shop = (rb if rb and (not yb or rb.total <= yb.total) else yb)

        values = [
            judgment,
            row["no"],
            p.title,
            p.amazon_price,
            rb.total if rb else "",
            yb.total if yb else "",
            best_shop.shop_name if best_shop else "",
            p.asin,
            p.jan,
            best_source,
            best_price,
            ref_fee,
            fba_fee,
            pr_r.profit if pr_r else "",
            pr_y.profit if pr_y else "",
            rb.url if rb else "",
            yb.url if yb else "",
        ]

        bg_color = JUDGMENT_COLORS.get(judgment, "FFFFFF")
        row_fill = PatternFill("solid", fgColor=bg_color)

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.fill = row_fill

        # URLをハイパーリンク化
        for col, url in [(16, rb.url if rb else None), (17, yb.url if yb else None)]:
            if url:
                cell = ws.cell(row=r, column=col)
                cell.hyperlink = url
                cell.font = LINK_FONT

    # 列幅設定
    col_widths = {
        1:  30, # 判定
        2:  6,  # No.
        3:  45, # 商品名
        4:  14, # Amazon価格
        5:  12, # 楽天最安値
        6:  12, # ヤフー最安値
        7:  20, # ショップ名
        8:  14, # ASIN
        9:  16, # EAN
        10: 12, # 最安仕入先
        11: 12, # 仕入最安値
        12: 14, # Amazon手数料
        13: 12, # FBA配送料
        14: 16, # 粗利(楽天)
        15: 16, # 粗利(ヤフー)
        16: 45, # 楽天URL
        17: 45, # ヤフーURL
    }
    for col, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    # 1行目の高さ
    ws.row_dimensions[1].height = 20


if __name__ == "__main__":
    main()
