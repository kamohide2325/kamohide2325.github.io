"""
千葉県築古戸建リサーチ結果（2026年4月28日）を
Excelに保存してGoogleスプレッドシートへアップロードする

実行方法:
    cd ~/kamohide2325.github.io/arbitrage
    python3 upload_research_0428.py
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from google_uploader import upload_to_sheets, open_url


def build_excel(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "千葉県築古戸建リサーチ0428"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    blue_fill    = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    green_fill   = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
    orange_fill  = PatternFill(start_color="FF8C00", end_color="FF8C00", fill_type="solid")
    red_fill     = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    lightgreen   = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    lightyellow  = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    lightorange  = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    lightred     = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")

    headers = [
        "No.", "所在地", "最寄り駅・アクセス", "価格(万円)", "表面利回り",
        "推定月額賃料", "土地面積", "建物面積", "築年（築年数）",
        "間取り", "駐車場", "構造", "ステータス", "URL", "短評・注意点",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = blue_fill
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[1].height = 30

    def section_row(row, label, fill):
        c = ws.cell(row=row, column=1, value=label)
        c.fill = fill
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
        ws.row_dimensions[row].height = 22
        for col in range(2, 16):
            ws.cell(row=row, column=col).fill = fill

    def data_row(row, cols_1_13, url, note, fill):
        for col, val in enumerate(cols_1_13, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.fill = fill
            c.font = Font(size=10)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c.border = border
        # 14列目：URL（クリッカブルリンク）
        uc = ws.cell(row=row, column=14, value="🔗 開く" if url else "—")
        if url:
            uc.hyperlink = url
            uc.font = Font(color="0070C0", underline="single", size=10)
        else:
            uc.font = Font(size=10, color="999999")
        uc.fill = fill
        uc.alignment = Alignment(horizontal="center", vertical="center")
        uc.border = border
        # 15列目：短評
        nc = ws.cell(row=row, column=15, value=note)
        nc.fill = fill
        nc.font = Font(size=10)
        nc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        nc.border = border
        ws.row_dimensions[row].height = 60

    # ===== 候補物件（個別URL確認済みのみ掲載） =====
    section_row(2, "【候補物件】個別URLが確認できた物件のみ掲載（★は特に注目）", green_fill)

    # row_data: (no, address, station, price, yield, rent, land, bldg, built, layout, parking, structure, status, note, url, fill)
    candidates = [
        (1, "千葉県茂原市押日", "JR外房線 新茂原駅\n徒歩44分（車必須）",
         190, "37.89%", "6.0万円", "204.55㎡", "62.7㎡", "1974年（築52年）",
         "3LDK", "あり（1台）", "木造2階", "空室・売主直接",
         "★最有力★ 190万・利回り38%・土地200㎡超・空室・売主直接と条件揃う。"
         "水道公営・下水は浄化槽（汲み取りでない）で除外対象外。"
         "新茂原駅44分は非現実的（車必須）。"
         "再建築可否・傾き・白アリは問合せ必須。",
         "https://www.kenbiya.com/pp8/s/chiba/mobara-shi/re_43697371kk/", lightgreen),
        (2, "千葉県茂原市（下矢返）", "JR外房線 茂原駅\nアクセス要確認",
         360, "14.00%", "4.2万円", "要確認", "要確認", "1983年（築43年）",
         "3LDK", "あり（2台）", "木造", "OC要確認",
         "360万・利回り14%・駐車場2台・3LDK。"
         "OC（オーナーチェンジ）の可能性あり→問合せで必ず確認。"
         "水道・下水・再建築可否・駐車場詳細も要問合せ。",
         "https://www.kenbiya.com/pp8/s/chiba/mobara-shi/re_4284952dvv/", lightyellow),
    ]

    r = 3
    for row_data in candidates:
        fill  = row_data[-1]
        url   = row_data[-2]
        note  = row_data[-3]
        cols  = list(row_data[:-3])  # no 〜 status（13列）
        data_row(r, cols, url, note, fill or PatternFill())
        r += 1

    # ===== URL未取得のため非掲載 =====
    r += 1
    section_row(r, "【URL未取得のため非掲載】個別ページURLが確認できなかった物件 → 各サイトで手動確認推奨", orange_fill)
    r += 1

    no_url_list = [
        ("—", "千葉市若葉区大宮町", "京成千原線 大森台駅 3.3km（車必須）", "220万", "21.81%", "4.0万円",
         "117.73㎡", "45.44㎡", "1971年（築55年）", "不明", "不明", "木造平屋", "URL未取得"),
        ("—", "千葉市若葉区大宮町", "京成千原線 大森台駅 37分徒歩", "250万", "17.76%", "3.7万円",
         "132.23㎡", "54.64㎡", "1974年（築52年）", "不明", "不明", "木造2階", "URL未取得"),
        ("—", "千葉市若葉区高根町", "千葉都市モノレール 千城台駅\nバス16分+徒歩5分", "250万", "26.40%", "5.5万円",
         "113.79㎡", "98.08㎡", "1980年（築46年）", "不明", "不明", "木造2階", "URL未取得"),
        ("—", "千葉県茂原市（鷲巣）", "要確認", "280万", "20.57%", "4.8万円",
         "69.27㎡", "49.65㎡", "1971年（築55年）", "不明", "不明", "木造平屋", "URL未取得"),
    ]
    for row_data in no_url_list:
        for col, val in enumerate(row_data, 1):
            c = ws.cell(row=r, column=col, value=val)
            c.fill = lightorange
            c.font = Font(size=10)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c.border = border
        uc = ws.cell(row=r, column=14, value="—")
        uc.fill = lightorange
        uc.font = Font(size=10, color="999999")
        uc.alignment = Alignment(horizontal="center", vertical="center")
        uc.border = border
        note_text = (
            "個別URLが健美家・楽待等で確認できなかったため非掲載。"
            "千葉市若葉区は3件とも健美家 千葉市若葉区ページで手動検索推奨。"
            if "千葉市" in row_data[1]
            else "茂原市の健美家ページで手動検索推奨。"
        )
        nc = ws.cell(row=r, column=15, value=note_text)
        nc.fill = lightorange
        nc.font = Font(size=10)
        nc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        nc.border = border
        ws.row_dimensions[r].height = 45
        r += 1

    # ===== 除外確定 =====
    r += 1
    section_row(r, "【除外確定】絶対除外条件に該当 / 予算超過", red_fill)
    r += 1

    excluded_list = [
        ("❌", "千葉県八街市八街い", "JR総武本線 八街駅 車15分", "180万", "33.33%", "5万円",
         "129㎡", "不明", "1975年（築51年）", "不明", "2台", "木造", "再建築不可",
         "再建築不可のため絶対除外"),
        ("❌", "千葉県茂原市七渡", "本納駅 徒歩35分（車必須）", "450万", "16.00%", "6万円",
         "147.39㎡", "103.12㎡", "1991年（築35年）", "4LDK", "1台", "木造2階", "予算超過",
         "450万は予算400万超過のため除外（個別URL: re_4286345p1s）"),
        ("❌", "千葉県鎌ヶ谷市", "東武野田線 馬込沢駅 徒歩19分", "450万", "15.20%", "5.7万円",
         "93.18㎡", "70.38㎡", "1976年（築50年）", "4DK", "不明", "木造2階", "予算超過",
         "450万は予算400万超過のため除外（個別URL: re_41599532wo）"),
        ("❌", "千葉県東金市", "—", "398万", "14.17%", "4.7万円",
         "不明", "不明", "不明", "不明", "不明", "不明", "オーナーチェンジ",
         "OC（賃貸中）のため絶対除外（個別URL: re_4266578par）"),
    ]
    for row_data in excluded_list:
        note = row_data[-1]
        cols = list(row_data[:-1])
        for col, val in enumerate(cols, 1):
            c = ws.cell(row=r, column=col, value=val)
            c.fill = lightred
            c.font = Font(size=10)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c.border = border
        uc = ws.cell(row=r, column=14, value="—")
        uc.fill = lightred
        uc.font = Font(size=10, color="999999")
        uc.alignment = Alignment(horizontal="center", vertical="center")
        uc.border = border
        nc = ws.cell(row=r, column=15, value=note)
        nc.fill = lightred
        nc.font = Font(size=10)
        nc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        nc.border = border
        ws.row_dimensions[r].height = 30
        r += 1

    # 列幅
    col_widths = [5, 22, 20, 10, 11, 11, 10, 10, 17, 8, 10, 16, 14, 10, 55]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ===== シート2：調査メモ =====
    ws2 = wb.create_sheet("調査メモ・凡例")
    notes = [
        ("■ 調査情報", ""),
        ("調査日",     "2026年4月28日"),
        ("調査サイト", "健美家（Google検索 site:kenbiya.com 経由）・楽待（一部）"),
        ("備考",       "各サイトへの直接アクセスが403のため検索キャッシュ経由で収集。"
                       "個別URLが確認できた物件のみ掲載（ポリシー遵守）。"),
        ("", ""),
        ("■ 今回の注目点", ""),
        ("",           "茂原市押日 190万・利回り37.89%（物件1）が最有力。"
                       "水道公営・下水浄化槽・空室・売主直接と条件が揃っています。"),
        ("",           "千葉市若葉区の3件（220万/250万/250万台）はURLが取得できず非掲載。"
                       "千葉市内で利回り20%超の物件が複数存在する可能性あり。"),
        ("", ""),
        ("■ 検索条件", ""),
        ("予算",         "400万円以下"),
        ("利回り",       "10%以上"),
        ("駐車場",       "あり"),
        ("対象エリア",   "千葉県の「市」のみ（町・村は除外）"),
        ("除外市",       "野田市・館山市・銚子市・勝浦市・鴨川市・印西市・南房総市・匝瑳市・いすみ市"),
        ("絶対除外条件", "OC / 再建築不可 / 井戸水 / 汲み取り / 傾き / 白アリ"),
        ("", ""),
        ("■ 色の意味", ""),
        ("緑（薄）",    "★最注目候補（物件1: 茂原市押日 190万・利回り38%）"),
        ("黄（薄）",    "候補物件（OC等の確認待ち）"),
        ("橙（薄）",    "URL未取得のため非掲載 → 手動確認が必要"),
        ("赤（薄）",    "除外確定（除外条件に該当 / 予算超過）"),
        ("", ""),
        ("■ 次のアクション（優先順）", ""),
        ("優先①", "物件1（茂原市押日 190万）→ 健美家 re_43697371kk で詳細確認・問合せ。"
                  "再建築可否・傾き・白アリ・下水詳細を必ず確認"),
        ("優先②", "物件2（茂原市 360万）→ OC非該当を確認後に問合せ。"
                  "健美家 re_4284952dvv"),
        ("優先③", "千葉市若葉区3件（URL未取得）→ 健美家 千葉市若葉区ページを直接確認。"
                  "https://www.kenbiya.com/pp8/s/chiba/chiba-shi/4/"),
        ("優先④", "茂原市鷲巣280万（URL未取得）→ 健美家 茂原市ページで直接検索"),
        ("全物件共通", "OC非該当・水道公営・下水汲み取りでない・再建築可・駐車場あり・"
                      "傾きなし・白アリなし を必ず問合せで確認"),
        ("", ""),
        ("■ 手動確認リンク", ""),
        ("健美家 茂原市 戸建", "https://www.kenbiya.com/pp8/s/chiba/mobara-shi/"),
        ("健美家 千葉市若葉区 戸建", "https://www.kenbiya.com/pp8/s/chiba/chiba-shi/4/"),
        ("楽待 千葉県 戸建 利回り10%以上", "https://www.rakumachi.jp/syuuekibukken/area/prefecture/dimAll/?pref[]=12&dim=1004&gross=10"),
    ]
    for i, (k, v) in enumerate(notes, 1):
        ck = ws2.cell(row=i, column=1, value=k)
        cv = ws2.cell(row=i, column=2, value=v)
        if k.startswith("■"):
            ck.fill = blue_fill
            cv.fill = blue_fill
            ck.font = Font(bold=True, color="FFFFFF", size=10)
            cv.font = Font(color="FFFFFF", size=10)
        elif k.startswith("優先") or k in ("全物件共通",):
            ck.font = Font(bold=True, size=10)
            cv.font = Font(size=10)
        else:
            ck.font = Font(size=10)
            cv.font = Font(size=10)
        cv.alignment = Alignment(wrap_text=True)
    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 80

    wb.save(path)
    print(f"Excel保存: {path}")


def main():
    desktop = Path.home() / "Desktop" / "千葉県築古戸建リサーチ結果_20260428.xlsx"
    build_excel(desktop)

    print("Googleスプレッドシートへアップロード中...")
    url = upload_to_sheets(str(desktop))
    print(f"アップロード完了: {url}")

    open_url(url)
    print("ブラウザで開きました。")


if __name__ == "__main__":
    main()
